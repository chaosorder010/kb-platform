from __future__ import annotations

import os
import re
import uuid
from pathlib import Path
from typing import Any, Callable

from knowledge.units.chunk_writer import ChunkWriter, MemoryChunkWriter
from knowledge.units.repository import KnowledgeRepository
from knowledge.units.schemas import utc_now_iso

ALLOWED_EXTENSIONS = {".pdf", ".md", ".markdown", ".txt", ".docx", ".doc"}
EXT_TO_TYPE = {
    ".pdf": "pdf",
    ".md": "md",
    ".markdown": "md",
    ".txt": "txt",
    ".docx": "docx",
    ".doc": "doc",
}


def detect_file_type(filename: str) -> str | None:
    ext = Path(filename).suffix.lower()
    return EXT_TO_TYPE.get(ext)


def title_from_filename(filename: str) -> str:
    stem = Path(filename).stem.strip()
    return stem or filename


class KnowledgeService:
    def __init__(
        self,
        repository: KnowledgeRepository,
        chunk_writer: ChunkWriter | None = None,
        pipeline: Callable[..., list[dict[str, Any]]] | None = None,
    ) -> None:
        self._repo = repository
        self._chunk_writer = chunk_writer or MemoryChunkWriter()
        self._pipeline = pipeline or self._default_light_pipeline
        self._code_seq = 0

    def create_unit_for_file(
        self,
        *,
        filename: str,
        content: bytes,
        creator_id: str,
        creator_name: str,
        category: str = "",
        status: str = "processing",
    ) -> dict[str, Any]:
        file_type = detect_file_type(filename)
        if file_type is None:
            raise ValueError("仅支持 PDF、Markdown、TXT、Word（.doc/.docx）格式")
        now = utc_now_iso()
        self._code_seq += 1
        unit_id = f"ku-{uuid.uuid4().hex[:12]}"
        unit = {
            "id": unit_id,
            "unit_code": f"KU-{now[:10].replace('-', '')}-{self._code_seq:04d}",
            "title": title_from_filename(filename),
            "content": "",
            "summary": "",
            "category": category,
            "source_file_name": filename,
            "file_type": file_type,
            "file_size": len(content),
            "status": status,
            "creator_id": creator_id,
            "creator_name": creator_name,
            "created_at": now,
            "updated_at": now,
            "data_permissions": [],
            "permission_summary": "无数据权限",
            "tags": [],
            "attachments": [],
            "raw_content": content,
        }
        return self._repo.insert_unit(unit)

    def import_files(
        self,
        *,
        files: list[tuple[str, bytes]],
        creator_id: str,
        creator_name: str,
        category: str = "",
    ) -> list[dict[str, str]]:
        if not files:
            raise ValueError("请至少上传一个文件")
        tasks: list[dict[str, str]] = []
        for filename, content in files:
            if detect_file_type(filename) is None:
                raise ValueError(f"不支持的文件格式: {filename}")
            unit = self.create_unit_for_file(
                filename=filename,
                content=content,
                creator_id=creator_id,
                creator_name=creator_name,
                category=category,
            )
            task_id = uuid.uuid4().hex[:12]
            self._repo.save_task(
                {
                    "task_id": task_id,
                    "unit_id": unit["id"],
                    "filename": filename,
                    "status": "pending",
                    "progress": 0,
                    "done_list": [],
                    "running_list": ["排队中"],
                    "error": "",
                    "content": content,
                }
            )
            tasks.append(
                {
                    "task_id": task_id,
                    "unit_id": unit["id"],
                    "filename": filename,
                }
            )
        return tasks

    def get_task_status(self, task_id: str) -> dict[str, Any] | None:
        return self._repo.get_task(task_id)

    def list_units(
        self,
        *,
        q: str | None = None,
        title: str | None = None,
        category: str | None = None,
        status: str | None = None,
        page: int = 1,
        page_size: int = 20,
    ) -> dict[str, Any]:
        items, total = self._repo.list_units(
            q=q,
            title=title,
            category=category,
            status=status,
            page=page,
            page_size=page_size,
        )
        public_items = []
        for item in items:
            public_items.append(self._to_public_unit(item))

        return {
            "items": public_items,
            "total": total,
            "page": page,
            "page_size": page_size,
        }

    def get_unit(self, unit_id: str) -> dict[str, Any] | None:
        unit = self._repo.get_unit(unit_id)
        if unit is None:
            return None
        public = self._to_public_unit(unit)
        public["attachments"] = self._repo.list_attachments(unit_id)
        return public

    def update_unit(
        self,
        unit_id: str,
        *,
        editor_id: str,
        editor_name: str = "",
        title: str | None = None,
        content: str | None = None,
        tags: list[str] | None = None,
        status: str | None = None,
        category: str | None = None,
        summary: str | None = None,
    ) -> dict[str, Any]:
        unit = self._repo.get_unit(unit_id)
        if unit is None:
            raise LookupError("知识单元不存在")

        current_status = unit.get("status") or "draft"
        next_status = status if status is not None else current_status
        if status is not None and not self._can_transition(current_status, next_status):
            raise ValueError(f"不允许从 {current_status} 转换到 {next_status}")

        patch: dict[str, Any] = {"updated_at": utc_now_iso()}
        if title is not None:
            patch["title"] = title
        if content is not None:
            patch["content"] = content
        if tags is not None:
            patch["tags"] = list(tags)
        if status is not None:
            patch["status"] = next_status
        if category is not None:
            patch["category"] = category
        if summary is not None:
            patch["summary"] = summary

        updated = self._repo.update_unit(unit_id, patch)
        if updated is None:
            raise LookupError("知识单元不存在")

        versions = self._repo.list_versions(unit_id)
        version_no = (versions[0]["version"] + 1) if versions else 1
        self._repo.save_version(
            {
                "id": f"ver-{uuid.uuid4().hex[:12]}",
                "unit_id": unit_id,
                "version": version_no,
                "title": updated.get("title") or "",
                "content": updated.get("content") or "",
                "tags": list(updated.get("tags") or []),
                "status": updated.get("status") or "",
                "editor_id": editor_id,
                "editor_name": editor_name,
                "created_at": utc_now_iso(),
            }
        )
        public = self._to_public_unit(updated)
        public["attachments"] = self._repo.list_attachments(unit_id)
        return public

    def list_versions(self, unit_id: str) -> list[dict[str, Any]]:
        if self._repo.get_unit(unit_id) is None:
            raise LookupError("知识单元不存在")
        return self._repo.list_versions(unit_id)

    def add_attachment(
        self,
        unit_id: str,
        *,
        filename: str,
        content: bytes,
        content_type: str = "",
        uploaded_by: str = "",
    ) -> dict[str, Any]:
        if self._repo.get_unit(unit_id) is None:
            raise LookupError("知识单元不存在")
        attachment_id = f"att-{uuid.uuid4().hex[:12]}"
        object_key = f"attachments/{unit_id}/{attachment_id}/{filename}"
        self._repo.put_object(object_key, content)
        attachment = {
            "id": attachment_id,
            "unit_id": unit_id,
            "filename": filename,
            "object_key": object_key,
            "size": len(content),
            "content_type": content_type,
            "uploaded_by": uploaded_by,
            "uploaded_at": utc_now_iso(),
        }
        saved = self._repo.save_attachment(attachment)
        attachments = self._repo.list_attachments(unit_id)
        self._repo.update_unit(
            unit_id,
            {"attachments": attachments, "updated_at": utc_now_iso()},
        )
        return {
            "id": saved["id"],
            "filename": saved["filename"],
            "object_key": saved["object_key"],
            "size": saved["size"],
            "content_type": saved.get("content_type") or "",
            "uploaded_by": saved.get("uploaded_by") or "",
            "uploaded_at": saved.get("uploaded_at") or "",
        }

    def set_permissions(
        self,
        unit_id: str,
        permissions: list[dict[str, Any]],
    ) -> dict[str, Any]:
        if self._repo.get_unit(unit_id) is None:
            raise LookupError("知识单元不存在")
        normalized: list[dict[str, Any]] = []
        for item in permissions:
            ptype = item.get("type")
            if ptype not in {"global", "department", "role", "user"}:
                raise ValueError(f"不支持的权限类型: {ptype}")
            pid = item.get("id") or ("*" if ptype == "global" else "")
            if ptype != "global" and not pid:
                raise ValueError("权限目标 id 不能为空")
            normalized.append(
                {
                    "type": ptype,
                    "id": pid if ptype != "global" else "*",
                    "name": item.get("name") or "",
                }
            )
        summary = self._permission_summary(normalized)
        updated = self._repo.update_unit(
            unit_id,
            {
                "data_permissions": normalized,
                "permission_summary": summary,
                "updated_at": utc_now_iso(),
            },
        )
        if updated is None:
            raise LookupError("知识单元不存在")
        public = self._to_public_unit(updated)
        public["attachments"] = self._repo.list_attachments(unit_id)
        return public

    def check_permissions(
        self,
        *,
        unit_ids: list[str],
        user_id: str,
        department_id: str = "",
        role_ids: list[str] | None = None,
    ) -> list[dict[str, Any]]:
        role_ids = role_ids or []
        results: list[dict[str, Any]] = []
        for unit_id in unit_ids:
            unit = self._repo.get_unit(unit_id)
            if unit is None:
                results.append({"unit_id": unit_id, "authorized": False})
                continue
            perms = unit.get("data_permissions") or []
            results.append(
                {
                    "unit_id": unit_id,
                    "authorized": self._is_authorized(
                        perms,
                        user_id=user_id,
                        department_id=department_id,
                        role_ids=role_ids,
                    ),
                }
            )
        return results

    def delete_units(self, unit_ids: list[str]) -> int:
        if not unit_ids:
            return 0
        return self._repo.delete_units(unit_ids)

    @staticmethod
    def _can_transition(current: str, target: str) -> bool:
        if current == target:
            return True
        allowed = {
            "draft": {"published", "disabled"},
            "published": {"draft", "disabled"},
            "disabled": {"draft", "published"},
            "failed": {"draft"},
            "processing": set(),
            "pending": set(),
            "completed": {"draft", "published", "disabled"},
        }
        return target in allowed.get(current, set())

    @staticmethod
    def _permission_summary(permissions: list[dict[str, Any]]) -> str:
        if not permissions:
            return "无数据权限"
        labels: list[str] = []
        for item in permissions:
            ptype = item.get("type")
            name = item.get("name") or item.get("id") or ""
            if ptype == "global":
                labels.append("全局")
            elif ptype == "department":
                labels.append(f"部门:{name}")
            elif ptype == "role":
                labels.append(f"角色:{name}")
            elif ptype == "user":
                labels.append(f"用户:{name}")
        return " / ".join(labels) if labels else "已配置权限"

    @staticmethod
    def _is_authorized(
        permissions: list[dict[str, Any]],
        *,
        user_id: str,
        department_id: str,
        role_ids: list[str],
    ) -> bool:
        if not permissions:
            return False
        for item in permissions:
            ptype = item.get("type")
            pid = item.get("id")
            if ptype == "global":
                return True
            if ptype == "user" and pid == user_id:
                return True
            if ptype == "department" and pid and pid == department_id:
                return True
            if ptype == "role" and pid in role_ids:
                return True
        return False

    def run_import_task_sync(self, task_id: str) -> list[dict[str, Any]]:
        task = self._repo.get_task(task_id)
        if task is None:
            raise ValueError("任务不存在")
        unit_id = task["unit_id"]
        filename = task["filename"]
        content: bytes = task.get("content") or b""

        steps = ["读取文件", "解析内容", "生成切片", "写入向量"]
        done: list[str] = []
        try:
            self._repo.update_task(
                task_id,
                {
                    "status": "processing",
                    "progress": 5,
                    "running_list": [steps[0]],
                    "done_list": [],
                },
            )
            text = self._extract_text(filename, content)
            done.append(steps[0])
            self._repo.update_task(
                task_id,
                {
                    "progress": 30,
                    "done_list": list(done),
                    "running_list": [steps[1]],
                },
            )
            done.append(steps[1])
            self._repo.update_task(
                task_id,
                {
                    "progress": 55,
                    "done_list": list(done),
                    "running_list": [steps[2]],
                },
            )
            chunks = self._pipeline(unit_id=unit_id, filename=filename, text=text)
            done.append(steps[2])
            self._repo.update_task(
                task_id,
                {
                    "progress": 80,
                    "done_list": list(done),
                    "running_list": [steps[3]],
                },
            )
            written = self._chunk_writer.write_chunks(unit_id, chunks)
            self._repo.save_chunks(unit_id, written)
            done.append(steps[3])
            now = utc_now_iso()
            self._repo.update_unit(
                unit_id,
                {
                    "content": text,
                    "summary": text[:120],
                    "status": "published",
                    "updated_at": now,
                },
            )
            self._repo.update_task(
                task_id,
                {
                    "status": "completed",
                    "progress": 100,
                    "done_list": list(done),
                    "running_list": [],
                    "error": "",
                },
            )
            return written
        except Exception as exc:
            self._repo.update_task(
                task_id,
                {
                    "status": "failed",
                    "progress": max(task.get("progress", 0), 10),
                    "running_list": [],
                    "done_list": done,
                    "error": str(exc),
                },
            )
            self._repo.update_unit(
                unit_id,
                {"status": "failed", "updated_at": utc_now_iso()},
            )
            raise

    def _to_public_unit(self, item: dict[str, Any]) -> dict[str, Any]:
        perms = item.get("data_permissions") or []
        attachments = item.get("attachments")
        if attachments is None:
            attachments = self._repo.list_attachments(item["id"])
        return {
            "id": item["id"],
            "unit_code": item["unit_code"],
            "title": item["title"],
            "category": item.get("category") or "",
            "file_type": item["file_type"],
            "source_file_name": item.get("source_file_name") or "",
            "permission_summary": item.get("permission_summary")
            or ("无数据权限" if not perms else "已配置权限"),
            "data_permissions": perms,
            "creator_id": item.get("creator_id") or "",
            "creator_name": item.get("creator_name") or "",
            "created_at": item.get("created_at") or "",
            "updated_at": item.get("updated_at") or "",
            "status": item.get("status") or "",
            "summary": item.get("summary") or "",
            "content": item.get("content") or "",
            "file_size": item.get("file_size") or 0,
            "tags": list(item.get("tags") or []),
            "attachments": list(attachments or []),
        }

    def _extract_text(self, filename: str, content: bytes) -> str:
        file_type = detect_file_type(filename)
        if file_type in {"md", "txt"}:
            return content.decode("utf-8", errors="ignore")
        if file_type == "pdf":
            text = content.decode("latin-1", errors="ignore")
            cleaned = re.sub(r"[^\x20-\x7E\u4e00-\u9fff\n]", " ", text)
            cleaned = cleaned.strip() or f"[PDF] {title_from_filename(filename)}"
            return cleaned
        if file_type == "docx":
            return self._extract_docx_text(content)
        if file_type == "doc":
            raise ValueError("不支持旧版 .doc 二进制格式，请另存为 .docx 后重试")
        raise ValueError("不支持的文件格式")

    @staticmethod
    def _extract_docx_text(content: bytes) -> str:
        from io import BytesIO

        from docx import Document

        try:
            document = Document(BytesIO(content))
        except Exception as exc:
            raise ValueError(f"Word 文档解析失败，请确认文件为有效的 .docx：{exc}") from exc

        parts: list[str] = []
        for paragraph in document.paragraphs:
            text = (paragraph.text or "").strip()
            if text:
                parts.append(text)
        for table in document.tables:
            for row in table.rows:
                cells = [(cell.text or "").strip() for cell in row.cells]
                cells = [c for c in cells if c]
                if cells:
                    parts.append(" | ".join(cells))
        text = "\n\n".join(parts).strip()
        if not text:
            raise ValueError("Word 文档解析结果为空")
        return text

    @staticmethod
    def _default_light_pipeline(
        *,
        unit_id: str,
        filename: str,
        text: str,
    ) -> list[dict[str, Any]]:
        pieces = [p.strip() for p in re.split(r"\n{2,}|(?<=。)", text) if p and p.strip()]
        if not pieces:
            pieces = [text or title_from_filename(filename)]
        chunks = []
        for idx, piece in enumerate(pieces):
            chunks.append(
                {
                    "content": piece,
                    "title": title_from_filename(filename),
                    "parent_title": "",
                    "file_title": title_from_filename(filename),
                    "item_name": "",
                    "unit_id": unit_id,
                    "chunk_index": idx,
                }
            )
        return chunks


def maybe_run_heavy_import_graph(
    task_id: str,
    import_file_path: str,
    file_dir: str,
    unit_id: str = "",
) -> None:
    _ = unit_id
    if os.getenv("KB_IMPORT_HEAVY", "").lower() not in {"1", "true", "yes"}:
        return
    from knowledge.service.upload_service import UpLoadService

    service = UpLoadService()
    service.run_import_graph(task_id, import_file_path, file_dir, unit_id=unit_id)
